from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SCR = r"C:/Users/gaboo/AppData/Local/Temp/claude/C--prueba-de-mudanzas/951e05ba-d604-4590-bb3c-401eb440a244/scratchpad/"
DFD = r"C:/Users/gaboo/Downloads/DFDs/"
OUT = r"C:/Users/gaboo/Downloads/Guia_Lectura_DFDs_Grupo5.docx"

NARANJA = RGBColor(0xC2, 0x55, 0x00)
GRIS = RGBColor(0x55, 0x55, 0x55)
HEAD_BG = "FCD9BD"
GROUP_BG = "EFEFEF"
ESTADO_BG = {"ok": "E3F1E6", "cambiar": "FFF3D1", "sobra": "FBE0DE"}

doc = Document()

# ---------- estilos base ----------
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for m in ("left_margin", "right_margin"):
    setattr(sec, m, Cm(2))
sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)

def set_font(style, size, bold=None, color=None):
    style.font.name = "Arial"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial") if style.element.rPr is not None and style.element.rPr.rFonts is not None else None
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color

normal = doc.styles["Normal"]
set_font(normal, 10.5)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.15
for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
    st = doc.styles[name]
    set_font(st, size, bold=True, color=NARANJA)
    st.paragraph_format.space_before = Pt(12 if name != "Heading 3" else 8)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.keep_with_next = True
    rf = st.element.rPr.rFonts
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rf.attrib.pop(qn(a), None)
for name in ("List Bullet", "List Number"):
    set_font(doc.styles[name], 10.5)
    doc.styles[name].paragraph_format.space_after = Pt(2)

# ---------- ayudas ----------
def para(text="", bold_prefix=None, size=None, italic=False, color=None, align=None, after=None):
    p = doc.add_paragraph()
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
        if size: r.font.size = Pt(size)
    r = p.add_run(text)
    r.italic = italic
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    if align: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p

def bullet(text, bold_prefix=None, style="List Bullet"):
    p = doc.add_paragraph(style=style)
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    p.add_run(text)
    return p

def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)

def cell_text(cell, text, bold=False, size=9, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    if color: r.font.color.rgb = color

def fixed_table(rows, widths_cm):
    t = doc.add_table(rows=rows, cols=len(widths_cm))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblPr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(w * 567)))
    for row in t.rows:
        for c, w in zip(row.cells, widths_cm):
            c.width = Cm(w)
    return t

def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    trPr.append(h)

def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    c = OxmlElement("w:cantSplit")
    c.set(qn("w:val"), "true")
    trPr.append(c)

def simple_table(header, rows, widths, size=9):
    t = fixed_table(1 + len(rows), widths)
    for c, h in zip(t.rows[0].cells, header):
        cell_text(c, h, bold=True, size=size)
        shade(c, HEAD_BG)
    repeat_header(t.rows[0])
    for r, data in zip(t.rows[1:], rows):
        no_split(r)
        for c, v in zip(r.cells, data):
            cell_text(c, v, size=size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

FLUJO_W = [0.9, 3.5, 3.0, 3.0, 6.6]

def flujos_table(grupos):
    total = sum(len(rows) + 1 for _, rows in grupos)
    t = fixed_table(1 + total, FLUJO_W)
    for c, h in zip(t.rows[0].cells, ["N°", "Flujo", "Sale de", "Llega a", "Observación"]):
        cell_text(c, h, bold=True)
        shade(c, HEAD_BG)
    repeat_header(t.rows[0])
    i, n = 1, 1
    for titulo, rows in grupos:
        gr = t.rows[i]
        no_split(gr)
        merged = gr.cells[0].merge(gr.cells[-1])
        cell_text(merged, titulo, bold=True, size=9, color=GRIS)
        merged.paragraphs[0].paragraph_format.keep_with_next = True
        shade(merged, GROUP_BG)
        i += 1
        for (flujo, de, a, estado, obs) in rows:
            r = t.rows[i]
            no_split(r)
            vals = [str(n), flujo, de, a, obs]
            for c, v in zip(r.cells, vals):
                cell_text(c, v, bold=(c is r.cells[1]))
            shade(r.cells[4], ESTADO_BG[estado])
            i += 1
            n += 1
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def preguntas(que, quien, como, almacenes):
    t = fixed_table(4, [3.6, 13.4])
    for r, (k, v) in zip(t.rows, [("¿Qué hace?", que), ("¿Quién lo estimula?", quien),
                                   ("¿Cómo?", como), ("¿Qué almacenes usa?", almacenes)]):
        no_split(r)
        cell_text(r.cells[0], k, bold=True, size=9.5)
        shade(r.cells[0], HEAD_BG)
        cell_text(r.cells[1], v, size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def imagen(path, pie):
    doc.add_picture(path, width=Cm(17))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = para(pie, size=8.5, italic=True, color=GRIS, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)

def falta(items):
    doc.add_heading("Lo que falta o hay que cambiar además", level=3)
    for it in items:
        bullet(it)

# ---------- pie de página ----------
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run("Grupo 5 · Guía de lectura de DFDs · Página ")
fr.font.size = Pt(8); fr.font.color.rgb = GRIS
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
rr = OxmlElement("w:r")
rpr = OxmlElement("w:rPr")
sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "16"); rpr.append(sz)
rr.append(rpr)
tt = OxmlElement("w:t"); tt.text = "1"; rr.append(tt)
fld.append(rr)
footer._p.append(fld)

# ================= PORTADA / INTRO =================
p = doc.add_paragraph()
r = p.add_run("Cómo leer los DFDs")
r.bold = True; r.font.size = Pt(24); r.font.color.rgb = NARANJA
p.paragraph_format.space_after = Pt(0)
para("Mudanzas y Fletes Salta · Grupo 5 · Guía de lectura flujo por flujo", size=11, color=GRIS, after=10)

para("La guía tiene dos partes. La primera explica cómo se lee cualquier DFD: qué es cada símbolo, "
     "cómo se lee cada flecha y cómo se nombra cada cosa. La segunda recorre los seis DFDs del grupo, "
     "flujo por flujo: cada flecha numerada, de dónde sale, quién la recibe y si está bien o qué hay que cambiar.")
para("Los números siguen el orden del dibujo, de arriba hacia abajo y agrupados por el elemento con el que "
     "conecta el proceso, para que se puedan encontrar mirando la imagen.")
para("Cuando un proceso consulta un almacén, se dibuja una flecha de ida con la clave de "
     "búsqueda y otra de vuelta con los datos encontrados.", bold_prefix="Notación. ")

# ================= PARTE 1 =================
doc.add_heading("Parte 1. Cómo se lee cualquier DFD", level=1)

doc.add_heading("1.1 Los cuatro símbolos", level=2)
simple_table(["Símbolo", "Cómo se dibuja", "Qué representa", "Ejemplo"], [
    ["Entidad externa", "Rectángulo cerrado", "Alguien de afuera del sistema que le manda datos o los recibe.", "Administrador"],
    ["Proceso", "Círculo u óvalo", "Algo que hace el sistema. Una sola tarea.", "Asignar Rol"],
    ["Almacén", "Rectángulo abierto a la derecha", "Datos guardados. Cada almacén es una tabla del DER.", "Usuarios"],
    ["Flujo", "Flecha con nombre", "Un dato que viaja de un lado a otro.", "DNI del usuario"],
], [3.0, 3.6, 7.2, 3.2])

doc.add_heading("1.2 Cómo se lee una flecha", level=2)
para("Toda flecha se lee con la misma frase:")
p = para("«[de dónde sale] le envía [nombre del flujo] a [quién lo recibe]».", align=WD_ALIGN_PARAGRAPH.CENTER)
p.runs[0].bold = True
simple_table(["Dirección", "Qué significa", "Cómo se dice"], [
    ["Entidad → Proceso", "La persona le da un dato al sistema.", "«El administrador le envía el usuario seleccionado al proceso.»"],
    ["Proceso → Entidad", "El sistema le muestra o le devuelve algo a la persona.", "«El proceso le envía la confirmación al administrador.»"],
    ["Proceso → Almacén", "Guarda un dato (registra) o manda la clave para buscar (consulta).", "«El proceso guarda el rol asignado en Usuarios.»"],
    ["Almacén → Proceso", "El proceso lee datos guardados. Es la respuesta a una consulta.", "«Usuarios le devuelve al proceso los datos del usuario.»"],
], [3.3, 6.2, 7.5])

doc.add_heading("1.3 Guardar o consultar: cómo se distinguen", level=2)
bullet("dos flechas. La de ida lleva la clave con la que se busca; la de vuelta trae lo que se encontró. "
       "Ejemplo: «DNI del usuario» va a Usuarios y vuelve «Datos del usuario».", bold_prefix="Consulta: ")
bullet("una sola flecha hacia el almacén, con el dato que se guarda. Ejemplo: «Rol asignado» va a Usuarios.",
       bold_prefix="Registro: ")
bullet("por el nombre y por si tiene vuelta. «DNI del usuario» tiene vuelta, así que es una clave. "
       "«Rol asignado» no tiene vuelta, así que se guarda.", bold_prefix="¿Cómo sé cuál es cuál? ")
bullet("tiene que ser un campo de esa tabla en el DER: DNI, patente, id, estado.", bold_prefix="La clave ")
bullet("si el proceso trae todo el listado (por ejemplo, todos los servicios), no hay clave que mandar: "
       "se dibuja solo la vuelta. Conviene confirmar con los ejemplos del profesor cómo lo dibuja él.",
       bold_prefix="Listados completos: ")

doc.add_heading("1.4 Lo que nunca puede aparecer", level=2)
bullet("Una entidad conectada directo a un almacén. Una persona nunca toca los datos: siempre pasa por un proceso.")
bullet("Flechas entre dos entidades o entre dos almacenes.")
bullet("Un almacén que confirma, avisa, valida o modifica. El almacén solo devuelve los datos que se le piden.")
bullet("Una vuelta sin ida: con esta notación, toda respuesta de un almacén necesita su clave (salvo los listados).")
bullet("Un dato que sale del proceso sin haber entrado: si el proceso entrega algo, lo tuvo que recibir o leer de algún lado.")
bullet("Flechas sueltas: cada flecha empieza y termina tocando un símbolo.")

doc.add_heading("1.5 Cómo se nombra cada cosa", level=2)
simple_table(["Elemento", "Cómo se escribe", "Bien", "Mal"], [
    ["Proceso", "Verbo en infinitivo + objeto. Una sola tarea.", "Asignar Rol, Registrar Pedido", "Gestionar Vehículo, Registro de pedido"],
    ["Flujo", "Sustantivo: el dato que viaja. Nunca un verbo.", "DNI del usuario, Rol nuevo", "modificar permiso, Solicitud Alta"],
    ["Almacén", "Sustantivo en plural, igual que la tabla del DER.", "Usuarios, Roles", "usuarios / Usurios, rol"],
    ["Entidad externa", "Sustantivo: quién es (su rol).", "Administrador, Cliente", "usuarios (si ya es un almacén)"],
], [2.8, 5.6, 4.2, 4.4])
para("Hay sustantivos que en realidad nombran una acción: «validación», «solicitud», «asignación», «modificación». "
     "Siguen contando lo que hace el proceso en vez de nombrar el dato. Por eso no sirven como flujos:",
     bold_prefix="Palabras trampa. ")
simple_table(["Cuenta la acción", "Nombra el dato"], [
    ["Validación Credenciales", "Datos de la cuenta"],
    ["Solicitud Alta", "Datos del vehículo"],
    ["Asignación de usuarios", "Empleados elegidos"],
    ["modificar permiso", "Permiso a modificar"],
], [8.5, 8.5])
para("¿Esto se puede guardar en una tabla o mostrar en una pantalla? «DNI del usuario» sí; «Validación» no. "
     "«Confirmación» se acepta en casi todos lados, pero un profesor estricto puede preferir el dato concreto "
     "(por ejemplo «Rol anterior y rol nuevo»).", bold_prefix="La prueba: ")
para("La misma palabra cambia de lugar según su forma: «Asignar» va en el proceso; «asignado» describe un dato "
     "y puede ir en un flujo («Rol asignado»); «asignación» no va en ningún lado.", bold_prefix="Ojo con las formas: ")

doc.add_heading("1.6 Cómo explicarlo en la defensa", level=2)
para("Las cuatro preguntas del profesor, siempre en este orden:")
for q in ["¿Qué hace? Una sola tarea.",
          "¿Quién lo estimula? La entidad que manda el primer dato.",
          "¿Cómo? Si registra (guarda en un almacén) o consulta (lee de un almacén), o las dos cosas.",
          "¿Qué almacenes usa?"]:
    bullet(q, style="List Number")
para("Para recorrer el diagrama: empezar por la entidad que lo estimula, seguir las flechas en el orden en que "
     "pasan las cosas y terminar con lo que se guarda y lo que se le devuelve a la entidad.")

# ================= PARTE 2 =================
doc.add_page_break()
doc.add_heading("Parte 2. Cada DFD, flujo por flujo", level=1)
para("En la columna Observación, el color indica el estado de cada flujo:")
leg = fixed_table(1, [5.66, 5.67, 5.67])
for c, (txt, k) in zip(leg.rows[0].cells, [("Bien", "ok"), ("Hay que cambiarle el nombre o el almacén", "cambiar"),
                                             ("Sobra o va al revés", "sobra")]):
    cell_text(c, txt, bold=True, size=9)
    shade(c, ESTADO_BG[k])
doc.add_paragraph()

# ---------- 2.1 Asignar Rol ----------
doc.add_heading("2.1 Asignar Rol (el modelo)", level=2)
para("Es el que armamos corregido, así que sirve de ejemplo de cómo tienen que quedar los demás. "
     "Archivo: DFDs\\DFD_Asignar_Rol.drawio.")
imagen(SCR + "asignar_rol.png", "DFD Asignar Rol")
preguntas("Le cambia el rol a un empleado.", "El Administrador.",
          "Consulta Usuarios y Roles, y registra el rol nuevo en Usuarios.", "Usuarios y Roles.")
flujos_table([
    ("Con el Administrador", [
        ("Usuario seleccionado", "Administrador", "Asignar Rol", "ok", "Bien. Estimula el proceso: dice a quién se le cambia el rol."),
        ("Rol nuevo", "Administrador", "Asignar Rol", "ok", "Bien. El rol que eligió el administrador."),
        ("Confirmación del cambio de rol", "Asignar Rol", "Administrador", "ok", "Bien. Versión más estricta: «Rol anterior y rol nuevo», que es lo que muestra la app."),
    ]),
    ("Con Usuarios", [
        ("DNI del usuario", "Asignar Rol", "Usuarios", "ok", "Bien. Ida de la consulta: la clave para buscar al usuario."),
        ("Datos del usuario", "Usuarios", "Asignar Rol", "ok", "Bien. Vuelta: rol actual y estado del usuario."),
        ("Rol asignado", "Asignar Rol", "Usuarios", "ok", "Bien. Registro: lo único que se guarda."),
    ]),
    ("Con Roles", [
        ("Id del rol nuevo", "Asignar Rol", "Roles", "ok", "Bien. Ida de la consulta: la clave para buscar el rol."),
        ("Datos del rol", "Roles", "Asignar Rol", "ok", "Bien. Vuelta: confirma que el rol existe."),
    ]),
])
doc.add_heading("Cómo se lee en voz alta", level=3)
para("«El administrador le envía al proceso el usuario seleccionado y el rol nuevo. El proceso le manda a Usuarios "
     "el DNI del usuario, y Usuarios le devuelve los datos del usuario. Después le manda a Roles el id del rol nuevo, "
     "y Roles le devuelve los datos del rol. Si el cambio es válido, el proceso guarda el rol asignado en Usuarios "
     "y le envía al administrador la confirmación del cambio de rol.»", italic=True)

# ---------- 2.2 Validar Usuario ----------
doc.add_page_break()
doc.add_heading("2.2 Validar Usuario", level=2)
imagen(DFD + "_DFDs_Mudanzas-VALIDAR_USUARIO.drawio.png", "DFD Validar Usuario, como está hoy")
preguntas("Controla que quien entra exista, tenga la clave correcta y esté activo, y le da acceso según su rol.",
          "El Usuario (cualquier rol).",
          "Consulta Usuarios, Roles y Permisos. Registra solo la fecha de último acceso.",
          "Usuarios, Roles y Permisos.")
flujos_table([
    ("Con la entidad Usuario", [
        ("Datos Credencial Usuario", "Usuario", "Validar Usuario", "ok", "Bien. El usuario y la clave que escribe la persona. Estimula el proceso."),
        ("Confirmacion Usuario", "Validar Usuario", "Usuario", "ok", "Bien."),
        ("Rol Usuario", "Validar Usuario", "Usuario", "ok", "Bien."),
        ("Permisos Rol", "Validar Usuario", "Usuario", "ok", "Bien."),
    ]),
    ("Con Usuarios", [
        ("Datos Credencial Usuario", "Validar Usuario", "Usuarios", "cambiar", "Ida de la consulta. Cambiar a «Nombre de usuario»: es la clave con la que se busca. La contraseña no sirve para buscar; se compara después, dentro del proceso."),
        ("Validacion Credenciales", "Usuarios", "Validar Usuario", "cambiar", "Vuelta. Cambiar a «Datos de la cuenta» (clave guardada, estado y rol). Validar lo hace el proceso, no el almacén."),
    ]),
    ("Con Permisos", [
        ("Permisos Rol", "Validar Usuario", "Permisos", "cambiar", "Ida. Cambiar a «Id del rol». En el DER, los permisos de cada rol están en PERMISOS_ROLES: conviene agregar ese almacén."),
        ("Permisos Rol", "Permisos", "Validar Usuario", "cambiar", "Vuelta. Cambiar a «Permisos del rol», para que no se llame igual que la ida."),
    ]),
    ("Con Roles", [
        ("Rol Usuario", "Validar Usuario", "Roles", "cambiar", "Ida. Cambiar a «Id del rol»."),
        ("Rol Usuario", "Roles", "Validar Usuario", "cambiar", "Vuelta. Cambiar a «Datos del rol»."),
        ("Validacion Credenciales", "Validar Usuario", "Roles", "sobra", "Sobra: el proceso no le manda nada de las credenciales a Roles."),
    ]),
])
falta(["«Fecha de último acceso»: Validar Usuario → Usuarios. Es lo único que este proceso guarda: "
       "la app registra la fecha cada vez que alguien entra."])

# ---------- 2.3 Gestionar Usuarios ----------
doc.add_page_break()
doc.add_heading("2.3 Gestionar Usuarios", level=2)
imagen(DFD + "_DFDs_Mudanzas-GESTIONAR USUARIOS.drawio.png", "DFD Gestionar Usuarios, como está hoy")
preguntas("Hoy hace alta, modificación y baja de usuarios, y además tiene el login mezclado. Son varias tareas.",
          "En el dibujo, el Usuario. Debería ser el Administrador.",
          "Registra en Usuarios (alta, correo y baja) y consulta Roles.",
          "Usuarios y Roles.")
flujos_table([
    ("Con la entidad Usuario", [
        ("Baja usuario", "Usuario", "Gestionar Usuario", "cambiar", "Cambiar a «DNI del usuario a dar de baja»: el dato que dice a quién."),
        ("Usuario nuevo", "Usuario", "Gestionar Usuario", "cambiar", "Cambiar a «Datos del usuario nuevo»."),
        ("Datos modificados", "Usuario", "Gestionar Usuario", "cambiar", "Cambiar a «Correo nuevo»: la consigna dice que solo se modifica el correo."),
        ("Datos Credencial Usuario", "Usuario", "Gestionar Usuario", "sobra", "Sobra: es del login (Validar Usuario). Además la flecha está suelta, no toca el rectángulo."),
        ("Confirmacion Usuario", "Gestionar Usuario", "Usuario", "sobra", "Sobra: es del login."),
        ("Rol Usuario", "Gestionar Usuario", "Usuario", "sobra", "Sobra: es del login."),
        ("Permisos Rol", "Gestionar Usuario", "Usuario", "sobra", "Sobra: es del login. Además sale sin que entre ningún dato de permisos al proceso."),
        ("Datos modificados", "Gestionar Usuario", "Usuario", "cambiar", "Cambiar a «Correo actualizado»."),
        ("Aviso Baja", "Gestionar Usuario", "Usuario", "ok", "Bien: el proceso le avisa al administrador."),
    ]),
    ("Con Usuarios", [
        ("Usuario Nuevo", "Gestionar Usuario", "Usuarios", "ok", "Bien. Registro del alta. Mejor: «Datos del usuario nuevo»."),
        ("Datos Credencial Usuario", "Gestionar Usuario", "Usuarios", "sobra", "Sobra: es del login."),
        ("Validacion Credenciales", "Usuarios", "Gestionar Usuario", "sobra", "Sobra: es del login."),
        ("Datos modificados", "Gestionar Usuario", "Usuarios", "cambiar", "Registro de la modificación. Cambiar a «Correo nuevo»."),
        ("Datos modificados", "Usuarios", "Gestionar Usuario", "cambiar", "Vuelta sin ida. Si se consulta al usuario antes de modificarlo: ida «DNI del usuario» y vuelta «Datos del usuario»."),
        ("Baja Usuario", "Gestionar Usuario", "Usuarios", "cambiar", "Registro de la baja. Cambiar a «Estado inactivo y fecha de baja»: la baja no borra, marca."),
        ("Aviso Baja", "Usuarios", "Gestionar Usuario", "sobra", "Sobra: un almacén no avisa."),
    ]),
    ("Con Roles", [
        ("Rol Usuario", "Gestionar Usuario", "Roles", "cambiar", "Ida (para el alta). Cambiar a «Id del rol»."),
        ("Rol Usuario", "Roles", "Gestionar Usuario", "cambiar", "Vuelta. Cambiar a «Datos del rol»."),
        ("Validacion Credenciales", "Gestionar Usuario", "Roles", "sobra", "Sobra: es del login."),
    ]),
])
falta([
    "La entidad tiene que ser Administrador, no Usuario: en la app solo el administrador puede hacer estas operaciones.",
    "Restablecer clave, que la consigna del Hito 2 pide: «DNI del usuario» (Administrador → proceso), "
    "«Clave temporal y cambio obligatorio» (proceso → Usuarios) y «Clave temporal» (proceso → Administrador).",
    "En el alta, la app le muestra al administrador el usuario que generó: «Nombre de usuario generado» (proceso → Administrador).",
    "Partirlo en cuatro procesos, uno por tarea: Registrar Usuario, Modificar Correo, Dar de baja Usuario y Restablecer Clave.",
])

# ---------- 2.4 Gestionar Vehículo ----------
doc.add_page_break()
doc.add_heading("2.4 Gestionar Vehículo", level=2)
imagen(DFD + "_DFDs_Mudanzas-GESTIONAR_VEHICULO.drawio.png", "DFD Gestionar Vehículo, como está hoy")
preguntas("Hoy hace alta, modificación, baja y cambio de estado de vehículos: cuatro tareas.",
          "El Administrador.", "Registra en Vehículos.", "Vehículos.")
flujos_table([
    ("Con el Administrador", [
        ("Solicitud Baja", "Administrador", "Gestionar Vehículo", "cambiar", "Cambiar a «Patente del vehículo»: el dato que dice cuál se da de baja."),
        ("Datos Vehiculo", "Administrador", "Gestionar Vehículo", "ok", "Bien. Los datos del alta."),
        ("Solicitud Alta", "Administrador", "Gestionar Vehículo", "sobra", "Sobra: los datos del alta ya viajan en el 2."),
        ("Estado Vehiculo", "Administrador", "Gestionar Vehículo", "ok", "Bien."),
        ("Datos modificacion", "Administrador", "Gestionar Vehículo", "ok", "Bien. Mejor: «Datos modificados del vehículo»."),
        ("Modificacion Vehiculo", "Gestionar Vehículo", "Administrador", "cambiar", "Cambiar a «Datos del vehículo actualizados»."),
        ("Confirmacion Vehiculo", "Gestionar Vehículo", "Administrador", "ok", "Bien."),
        ("Confirmacion Baja", "Gestionar Vehículo", "Administrador", "ok", "Bien."),
    ]),
    ("Con Vehículos", [
        ("Solicitud Baja", "Gestionar Vehículo", "Vehículos", "sobra", "Sobra: en la app, dar de baja es guardar el estado «Fuera de Servicio», que ya es el 12."),
        ("Solicitud Alta", "Gestionar Vehículo", "Vehículos", "sobra", "Sobra: no es un dato ni una clave."),
        ("Datos Vehiculo", "Gestionar Vehículo", "Vehículos", "ok", "Bien. Registro del alta."),
        ("Estado Vehiculo", "Gestionar Vehículo", "Vehículos", "ok", "Bien. Registro del estado (la baja también es esto)."),
        ("Datos modificacion", "Gestionar Vehículo", "Vehículos", "ok", "Bien. Registro de la modificación."),
        ("Modificacion Vehiculo", "Vehículos", "Gestionar Vehículo", "sobra", "Sobra: un almacén no modifica."),
        ("Confirmacion Vehiculo", "Vehículos", "Gestionar Vehículo", "sobra", "Sobra: un almacén no confirma."),
        ("Confirmacion Baja", "Vehículos", "Gestionar Vehículo", "sobra", "Sobra: un almacén no confirma."),
    ]),
])
falta([
    "Para controlar que la patente no esté repetida, o mostrar el vehículo antes de modificarlo: "
    "ida «Patente» (proceso → Vehículos) y vuelta «Datos del vehículo» (Vehículos → proceso).",
    "Partirlo en tres procesos: Registrar Vehículo, Modificar Vehículo y Cambiar Estado del Vehículo "
    "(la baja es uno de los estados).",
])

# ---------- 2.5 Registrar un Pedido ----------
doc.add_page_break()
doc.add_heading("2.5 Registrar un Pedido", level=2)
imagen(DFD + "_DFDs_Mudanzas-REGISTRAR_UN_PEDIDO.drawio.png", "DFD Registrar un Pedido, como está hoy")
preguntas("Registra el pedido de un cliente: fecha, servicio y vehículo.",
          "El Usuario del dibujo, que en realidad es el Cliente.",
          "Consulta Servicios, Vehículos y Usuarios. Registra en Pedidos (y en Detalle_Pedidos, que falta).",
          "Pedidos, Vehículos, Usuarios y Servicios.")
flujos_table([
    ("Con la entidad Usuario (el cliente)", [
        ("Disponibilidad de servicio", "Usuario", "Registrar un Pedido", "cambiar", "Cambiar a «Fecha del pedido»: es el dato que manda el cliente, y PEDIDOS lo necesita."),
        ("Servicios disponibles", "Registrar un Pedido", "Usuario", "ok", "Bien."),
        ("Selección de servicio", "Usuario", "Registrar un Pedido", "ok", "Bien. Mejor: «Servicio elegido»."),
        ("Disponibilidad de vehículo", "Usuario", "Registrar un Pedido", "sobra", "Sobra: con la fecha del 1 alcanza."),
        ("Vehículo disponibles", "Registrar un Pedido", "Usuario", "ok", "Bien. Corregir a «Vehículos disponibles»."),
        ("Seleccion de vehiculo", "Usuario", "Registrar un Pedido", "ok", "Bien. Mejor: «Vehículo elegido»."),
        ("Confirmacion pedido", "Registrar un Pedido", "Usuario", "ok", "Bien."),
    ]),
    ("Con el Administrador", [
        ("Confirmacion pedido", "Administrador", "Registrar un Pedido", "sobra", "Confirmar es otra tarea (Confirmar Pedido), y PEDIDOS no tiene columna de estado para guardarlo. Sacarlo o hacer un proceso aparte."),
        ("datos pedido", "Registrar un Pedido", "Administrador", "sobra", "Va junto con el 8."),
    ]),
    ("Con Pedidos", [
        ("Selección de servicio", "Registrar un Pedido", "Pedidos", "cambiar", "Registro, pero en otro almacén: según el DER, el servicio elegido se guarda en DETALLE_PEDIDOS."),
        ("Selección de vehiculo", "Registrar un Pedido", "Pedidos", "ok", "Bien. Registro de la patente. Mejor: «Patente elegida»."),
        ("Datos usuario", "Registrar un Pedido", "Pedidos", "ok", "Bien. Registro del cliente. Mejor: «Id del cliente»."),
        ("Datos de pedido", "Pedidos", "Registrar un Pedido", "sobra", "Sobra: vuelta sin ida. Para registrar no hace falta leer Pedidos."),
    ]),
    ("Con Vehículos", [
        ("Disponibilidad de vehículo", "Registrar un Pedido", "Vehículos", "cambiar", "Ida. Cambiar a la clave: «Estado buscado» (Disponible)."),
        ("Vehículo disponible", "Vehículos", "Registrar un Pedido", "ok", "Bien. Vuelta."),
        ("Selección Vehículo", "Registrar un Pedido", "Vehículos", "sobra", "Sobra: el vehículo elegido se guarda en Pedidos (11)."),
    ]),
    ("Con Usurios", [
        ("Datos usuario", "Registrar un Pedido", "Usurios", "cambiar", "Ida. Cambiar a «DNI del cliente» (el del cliente que inició sesión)."),
        ("Datos usuario", "Usurios", "Registrar un Pedido", "ok", "Bien. Vuelta. Mejor: «Datos del cliente»."),
    ]),
    ("Con Servicios", [
        ("Disponibilidad de servicio", "Registrar un Pedido", "Servicios", "sobra", "Sobra: para traer todos los servicios no hay clave; alcanza con la vuelta (20)."),
        ("Servicios disponibles", "Servicios", "Registrar un Pedido", "ok", "Bien. Vuelta con el listado."),
        ("seleccion de servicio", "Registrar un Pedido", "Servicios", "sobra", "Sobra: el servicio elegido se guarda en DETALLE_PEDIDOS (ver 10), no en Servicios."),
    ]),
])
falta([
    "Guardar la «Fecha del pedido» en Pedidos (proceso → Pedidos).",
    "Agregar el almacén Detalle_Pedidos, donde se guarda el servicio elegido (ver 10).",
    "Corregir el almacén «Usurios» a «Usuarios», y la entidad «Usuario» a «Cliente».",
    "Opcional: acá se podría cargar INVENTARIO_DETALLES (peso, cantidad, frágil) con lo que el cliente declara que va a mudar.",
])

# ---------- 2.6 Asignar recursos al pedido ----------
doc.add_page_break()
doc.add_heading("2.6 Asignar recursos al pedido", level=2)
imagen(DFD + "_DFDs_Mudanzas-ASIGNACION_DE_RECURSOS.drawio.png", "DFD Asignar recursos al pedido, como está hoy")
para("En este diagrama «usuarios» aparece dos veces: a la izquierda como entidad (los empleados que reciben la "
     "asignación) y arriba a la derecha como almacén. En la tabla se distinguen como «usuarios (entidad)» y "
     "«usuarios (almacén)».", size=9.5)
preguntas("Hoy asigna empleados, vehículo y materiales a un pedido: tres tareas.",
          "Los encargados.",
          "Consulta pedido, usuarios, servicio, vehiculo y recursos asignados. Debería registrar la asignación, pero hoy no la guarda en ningún lado.",
          "usuarios, servicio, pedido, vehiculo y recursos asignados.")
flujos_table([
    ("Con encargados", [
        ("datos de pedido(agenda)", "encargados", "Asignar recursos", "ok", "Bien. Mejor: «Pedido elegido» (el que se va a armar)."),
        ("disponibilidad de empleados", "encargados", "Asignar recursos", "sobra", "Sobra: el encargado no manda la disponibilidad; el proceso la busca en usuarios (12 y 14)."),
        ("Usuarios disponible", "Asignar recursos", "encargados", "ok", "Bien. Mejor: «Empleados disponibles»."),
        ("asignacion de usuarios", "encargados", "Asignar recursos", "cambiar", "Cambiar a «Empleados elegidos»: «asignación» nombra la acción, no el dato."),
        ("datos de pedido(agenda)", "Asignar recursos", "encargados", "ok", "Bien. Mejor: «Agenda de pedidos»."),
    ]),
    ("Con usuarios (entidad)", [
        ("asignacion de usuarios", "Asignar recursos", "usuarios (entidad)", "cambiar", "Cambiar a «Pedido asignado»: lo que recibe el empleado."),
        ("dato de servicio", "Asignar recursos", "usuarios (entidad)", "ok", "Bien."),
        ("estado de usuario", "usuarios (entidad)", "Asignar recursos", "cambiar", "Revisar: el proceso lo recibe pero no lo usa ni lo guarda en ningún lado."),
        ("dato de pedido", "Asignar recursos", "usuarios (entidad)", "ok", "Bien."),
        ("materiales asignados", "Asignar recursos", "usuarios (entidad)", "ok", "Bien."),
        ("datos vehiculos", "Asignar recursos", "usuarios (entidad)", "ok", "Bien. Mejor: «Vehículo asignado»."),
    ]),
    ("Con usuarios (almacén)", [
        ("disponibilidad de usuarios", "Asignar recursos", "usuarios (almacén)", "cambiar", "Ida. Cambiar a la clave: «Rol requerido» (chofer o ayudante)."),
        ("Estado de usuarios", "Asignar recursos", "usuarios (almacén)", "sobra", "La punta se confunde con la del 12. Si guarda «ocupado», USUARIOS no tiene ese campo: sobra."),
        ("Usuarios disponibles", "usuarios (almacén)", "Asignar recursos", "ok", "Bien. Vuelta. Mejor: «Empleados disponibles»."),
    ]),
    ("Con servicio", [
        ("dato de servicio", "servicio", "Asignar recursos", "ok", "Bien. Vuelta. Mejor: «Datos del servicio»."),
        ("datos de servicio", "Asignar recursos", "servicio", "cambiar", "Ida. Cambiar a la clave: «Id del servicio». Asignar no modifica el catálogo."),
    ]),
    ("Con pedido", [
        ("datos de pedido(agenda)", "pedido", "Asignar recursos", "ok", "Bien. Vuelta. Mejor: «Datos del pedido»."),
        ("datos de pedido(agenda)", "Asignar recursos", "pedido", "cambiar", "Ida. Cambiar a la clave: «Id del pedido», para que no se llame igual que la vuelta."),
    ]),
    ("Con vehiculo", [
        ("datos de vehiculo", "Asignar recursos", "vehiculo", "cambiar", "Ida. Cambiar a la clave: «Estado buscado» (Disponible)."),
        ("datos de veihculo", "vehiculo", "Asignar recursos", "ok", "Bien. Vuelta. Corregir «veihculo»."),
    ]),
    ("Con recursos asignados", [
        ("disponibilidad materialeas", "Asignar recursos", "recursos asignados", "sobra", "Sobra: la clave ya viaja en el 23. Además tiene error de tipeo."),
        ("disponibilidad de materiales", "recursos asignados", "Asignar recursos", "ok", "Bien. Vuelta. Mejor: «Materiales del servicio»."),
        ("seleccion de servicio", "Asignar recursos", "recursos asignados", "cambiar", "Ida. Cambiar a la clave: «Id del servicio»."),
        ("materiales asignados", "recursos asignados", "Asignar recursos", "sobra", "Va al revés: lo que se asigna se guarda (proceso → recursos asignados)."),
    ]),
])
falta([
    "Guardar los empleados asignados. Hoy «asignacion de usuarios» solo llega a los empleados y no queda en ningún "
    "almacén. En el DER tampoco hay dónde: RECURSOS_ASIGNADOS necesita id_pedido e id_usuario.",
    "Guardar el vehículo asignado: «Patente asignada» (proceso → pedido), porque la patente del pedido está en PEDIDOS.",
    "Renombrar la entidad «usuarios» a «Empleado», para que no se llame igual que el almacén, y «encargados» "
    "a «Administrador» si es él quien asigna.",
    "Partirlo en tres procesos: Asignar Empleados, Asignar Vehículo y Asignar Materiales.",
])

# ---------- Nombres de almacenes ----------
doc.add_heading("Nombres de almacenes en todos los DFDs", level=2)
para("El mismo almacén tiene que llamarse igual en todos los diagramas, y como la tabla del DER:")
simple_table(["Tabla del DER", "Cómo aparece hoy", "Cómo tiene que decir"], [
    ["USUARIOS", "Usuarios / usuarios / Usurios", "Usuarios"],
    ["VEHICULOS", "Vehículos / vehiculo", "Vehículos"],
    ["PEDIDOS", "Pedidos / pedido", "Pedidos"],
    ["SERVICIOS", "Servicios / servicio", "Servicios"],
    ["ROLES", "Roles", "Roles"],
    ["PERMISOS", "Permisos", "Permisos"],
    ["RECURSOS_ASIGNADOS", "recursos asignados", "Recursos Asignados"],
    ["DETALLE_PEDIDOS", "(no aparece)", "Detalle Pedidos"],
], [5.0, 6.5, 5.5])

doc.save(OUT)
print("OK", OUT)

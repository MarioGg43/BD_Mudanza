from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re, pathlib

BASE = pathlib.Path(__file__).parent
OUT = r"C:/Users/gaboo/Downloads/Guia_App_Movil_Figma_Grupo5.docx"
LINK = "https://claude.ai/artifact/9vJxwT88TZboVHaoA5akPq"
NARANJA = RGBColor(0xC2, 0x55, 0x00)
GRIS = RGBColor(0x55, 0x55, 0x55)
HEAD_BG = "FCD9BD"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2)
sec.top_margin = sec.bottom_margin = Cm(1.8)

def fuente(st, size, bold=None, color=None):
    st.font.name = "Arial"
    st.font.size = Pt(size)
    if bold is not None: st.font.bold = bold
    if color is not None: st.font.color.rgb = color
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"): rf.set(qn(a), "Arial")
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"): rf.attrib.pop(qn(a), None)

fuente(doc.styles["Normal"], 10.5)
doc.styles["Normal"].paragraph_format.space_after = Pt(4)
doc.styles["Normal"].paragraph_format.line_spacing = 1.15
for n, s in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
    st = doc.styles[n]; fuente(st, s, True, NARANJA)
    st.paragraph_format.space_before = Pt(12 if n != "Heading 3" else 8)
    st.paragraph_format.space_after = Pt(4); st.paragraph_format.keep_with_next = True
for n in ("List Bullet", "List Number"):
    fuente(doc.styles[n], 10.5); doc.styles[n].paragraph_format.space_after = Pt(2)

def para(texto="", negrita=None, size=None, italic=False, color=None, align=None, after=None):
    p = doc.add_paragraph()
    if negrita:
        r = p.add_run(negrita); r.bold = True
        if size: r.font.size = Pt(size)
    r = p.add_run(texto); r.italic = italic
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    if align is not None: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p

def item(texto, negrita=None, estilo="List Bullet"):
    p = doc.add_paragraph(style=estilo)
    if negrita: p.add_run(negrita).bold = True
    p.add_run(texto)
    return p

def numerar(lista):
    for i, t in enumerate(lista, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(str(i) + ".  "); r.bold = True; r.font.color.rgb = NARANJA
        p.add_run(t)

def sombra(celda, color):
    tcPr = celda._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color); tcPr.append(shd)

def texto_celda(c, t, bold=False, size=9):
    c.text = ""; p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
    r = p.add_run(t); r.bold = bold; r.font.size = Pt(size)

def tabla(encabezado, filas, anchos, size=9):
    t = doc.add_table(rows=1 + len(filas), cols=len(anchos))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); t._tbl.tblPr.append(lay)
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), anchos): gc.set(qn("w:w"), str(int(w * 567)))
    for fila in t.rows:
        for c, w in zip(fila.cells, anchos): c.width = Cm(w)
        trPr = fila._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trPr.append(cs)
    for c, h in zip(t.rows[0].cells, encabezado): texto_celda(c, h, True, size); sombra(c, HEAD_BG)
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); t.rows[0]._tr.get_or_add_trPr().append(th)
    for fila, datos in zip(t.rows[1:], filas):
        for c, v in zip(fila.cells, datos): texto_celda(c, v, False, size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

# pie de página
pie = sec.footer.paragraphs[0]; pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = pie.add_run("Grupo 5 · App móvil Mudanzas Salta · Página "); r.font.size = Pt(8); r.font.color.rgb = GRIS
fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
rr = OxmlElement("w:r"); rpr = OxmlElement("w:rPr"); sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "16"); rpr.append(sz); rr.append(rpr)
tt = OxmlElement("w:t"); tt.text = "1"; rr.append(tt); fld.append(rr); pie._p.append(fld)

# ================= portada =================
p = doc.add_paragraph(); r = p.add_run("App móvil Mudanzas Salta"); r.bold = True; r.font.size = Pt(24); r.font.color.rgb = NARANJA
p.paragraph_format.space_after = Pt(0)
para("Cómo probarla, cómo pasarla a Figma y qué muestra cada pantalla · Grupo 5 · Desarrollo Móvil", size=11, color=GRIS, after=10)
doc.add_picture(str(BASE / "consigna.png"), width=Cm(17))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
para("Las cinco vistas que pide la consigna, dentro de la app completa.", size=8.5, italic=True, color=GRIS, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)

doc.add_heading("1. Qué tienen", level=1)
tabla(["Qué", "Dónde", "Para qué"], [
    ["La app funcionando", LINK, "Probarla y presentarla. Se abre en la compu o en el celular."],
    ["Copia de la app", "Descargas\\App Mudanzas Movil\\index.html", "Abrirla con Chrome sin depender del link. Es la que se importa a Figma."],
    ["Capturas", "Descargas\\App Mudanzas Movil\\Pantallas (35 imágenes)", "Plan B para Figma y para el documento de entrega."],
], [3.6, 6.6, 6.8])
para("El link es privado: solo lo abre quien lo creó. Para que lo vean el grupo o el profesor, hay que compartirlo desde el botón Compartir de la página.", negrita="Ojo: ")
para("La consigna nombra Visio y Visily como herramientas. Figma no aparece, así que conviene confirmar con el profesor que lo acepta.", negrita="Antes de entregar: ")

# ================= probar =================
doc.add_heading("2. Cómo probar la app", level=1)
item("en la compu se ve dentro de un teléfono, con un panel a la derecha para entrar con cada rol sin escribir nada.", "En la compu: ")
item("abrir el link en el navegador del celular: ocupa toda la pantalla, como una app.", "En el celular: ")
item("los cambios quedan guardados en ese navegador. Antes de presentar, tocar \"Reiniciar demo\" (panel de la derecha o Perfil) para volver a los datos de ejemplo.", "Datos: ")
tabla(["Rol", "Usuario", "Contraseña", "Qué ve"], [
    ["Administrador", "sarmientoc", "Admin123!", "Panel, pedidos, asignación, cobros, vehículos, usuarios y servicios"],
    ["Chofer", "fernandezl", "Usuario123!", "Sus mudanzas; las inicia y las finaliza"],
    ["Ayudante", "romerol", "Temporal123!", "Primero le pide cambiar la clave (Hito 2)"],
    ["Cliente", "gomezm", "Cliente123!", "Pide mudanzas y sigue sus pedidos"],
], [2.6, 2.6, 2.6, 9.2])

doc.add_heading("Recorrido sugerido para la presentación", level=2)
numerar([
    "Inicio → Comenzar → Iniciar sesión. Mostrar que valida: una clave mal escrita da error.",
    "Entrar como María (cliente) → Pedir una mudanza: fecha y direcciones, servicios, vehículo, qué se lleva y resumen → Confirmar pedido.",
    "Entrar como Carlos (administrador) → Panel → \"Pedidos por confirmar\" → el pedido nuevo → Confirmar → Asignar equipo → Guardar asignación → Registrar cobro.",
    "Entrar como Luciana (chofer) → Próximas → la mudanza → Iniciar mudanza → Finalizar mudanza.",
    "Volver a María → Mis pedidos → Terminados: el pedido aparece Finalizado, con el equipo y el pago.",
    "Como Carlos: Vehículos (el CRUD de la consigna): buscar, ver, agregar, editar y dar de baja con confirmación.",
    "Como Carlos: Usuarios (Hito 2): alta con usuario automático, cambiar rol y restablecer clave. Intentar darse de baja a sí mismo muestra el candado.",
    "Perfil: editar datos, guardar y cerrar sesión.",
])

# ================= figma =================
doc.add_heading("3. Pasarla a Figma con html.to.design", level=1)
para("html.to.design es un plugin de Figma con una extensión para Chrome. Convierte una página web en capas editables de Figma: "
     "textos, botones, colores y formas se pueden modificar. El plan gratis permite hasta 10 importaciones cada 30 días e incluye la extensión. "
     "Como la página \"Todas las pantallas\" tiene las 35 vistas juntas, alcanza con una sola importación.")
numerar([
    "En Figma, abrir un archivo de diseño nuevo, ir a Plugins, buscar \"html.to.design\" y abrirlo. La primera vez pide iniciar sesión.",
    "En Chrome, instalar la extensión html.to.design desde la Chrome Web Store.",
    "Como la app es un archivo de la compu, hay que darle permiso a la extensión: abrir chrome://extensions, entrar a Detalles de html.to.design y activar \"Permitir el acceso a las URL de archivos\".",
    "Abrir con Chrome el archivo Descargas\\App Mudanzas Movil\\index.html y tocar \"Ver todas las pantallas\".",
    "Tocar el ícono de la extensión y capturar la página completa. Elegir \"Send to Figma plugin\" para mandarla directo, o guardar el archivo .h2d y arrastrarlo dentro del plugin.",
    "En Figma aparecen las 35 pantallas como capas editables; debajo de cada una queda su número y su nombre.",
])
para("Si la importación falla o se terminan las importaciones gratis, arrastren a Figma las imágenes de la carpeta Pantallas. "
     "Quedan como imágenes, no editables, pero sirven igual para armar el prototipo navegable.", negrita="Plan B: ")

doc.add_heading("4. Armar el prototipo en Figma", level=1)
numerar([
    "Cada pantalla tiene que ser un frame de 390 × 844, el tamaño de un iPhone. Si alguna quedó como grupo, seleccionarla y usar Ctrl + Alt + G (\"Frame selection\").",
    "Pasar a la pestaña Prototype, arriba a la derecha.",
    "Seleccionar un botón, arrastrar el círculo que aparece al costado hasta la pantalla de destino y elegir: On tap → Navigate to. La animación puede ser Instant o Smart animate.",
    "Para las flechas de volver, usar la acción Back.",
    "Los diálogos (confirmar baja, iniciar mudanza) se conectan con Open overlay, posición Center, con fondo oscuro.",
    "Marcar la pantalla Inicio como punto de partida (Flow starting point).",
    "Probarlo con el botón Present (▶). En el celular se ve con la app de Figma.",
])

doc.add_heading("Conexiones principales", level=2)
para("Los números son los de la página \"Todas las pantallas\" y de las imágenes de la carpeta Pantallas.", size=9.5, color=GRIS).paragraph_format.keep_with_next = True
tabla(["Desde", "Tocar", "Va a", "Tipo"], [
    ["1 Inicio", "Comenzar", "2 Iniciar sesión", "Navigate to"],
    ["2 Iniciar sesión", "Registrate", "3 Registrarse", "Navigate to"],
    ["2 Iniciar sesión", "¿Olvidaste tu contraseña?", "4 Recuperar contraseña", "Navigate to"],
    ["4 Recuperar contraseña", "Enviar contraseña temporal", "5 Contraseña enviada", "Navigate to"],
    ["2 Iniciar sesión", "Entrar rápido: Cliente", "7 Inicio del cliente", "Navigate to"],
    ["2 Iniciar sesión", "Entrar rápido: Administrador", "16 Panel", "Navigate to"],
    ["2 Iniciar sesión", "Entrar rápido: Chofer", "32 Mis mudanzas", "Navigate to"],
    ["2 Iniciar sesión", "Entrar rápido: Ayudante", "6 Cambio obligatorio de clave", "Navigate to"],
    ["7 Inicio del cliente", "Pedir una mudanza", "8 → 9 → 10 → 11 → 12 (Siguiente)", "Navigate to"],
    ["12 Paso 5", "Confirmar pedido", "13 Pedido enviado", "Navigate to"],
    ["14 Mis pedidos", "Una tarjeta", "15 Detalle del pedido", "Navigate to"],
    ["16 Panel", "Pedidos por confirmar", "17 Pedidos", "Navigate to"],
    ["17 Pedidos", "Pedido #1025", "18 Detalle del pedido", "Navigate to"],
    ["18 Detalle del pedido", "Rechazar", "19 Rechazar pedido", "Open overlay"],
    ["16 Panel · barra de abajo", "Vehículos", "22 Vehículos", "Navigate to"],
    ["22 Vehículos", "Ver", "23 Detalle del vehículo", "Navigate to"],
    ["22 Vehículos", "Botón +", "24 Agregar vehículo", "Navigate to"],
    ["22 Vehículos", "Editar", "25 Editar vehículo", "Navigate to"],
    ["22 Vehículos", "Baja", "26 Confirmar baja", "Open overlay"],
    ["16 Panel · barra de abajo", "Usuarios", "27 Usuarios", "Navigate to"],
    ["27 Usuarios", "Botón +", "28 Nuevo usuario", "Navigate to"],
    ["27 Usuarios", "Lucas Romero", "29 Detalle del usuario", "Navigate to"],
    ["29 Detalle del usuario", "Cambiar rol", "30 Cambiar rol", "Open overlay"],
    ["32 Mis mudanzas", "Mudanza #1023", "33 Detalle de la mudanza", "Navigate to"],
    ["33 Detalle de la mudanza", "Iniciar mudanza", "34 Iniciar mudanza", "Open overlay"],
    ["Cualquier barra de abajo", "Perfil", "35 Perfil", "Navigate to"],
    ["35 Perfil", "Cerrar sesión", "2 Iniciar sesión", "Navigate to"],
], [4.0, 4.4, 5.6, 3.0], size=8.5)

# ================= consigna =================
doc.add_heading("5. Qué pantalla cubre cada punto de la consigna", level=1)
tabla(["Vista de la consigna", "Pedido", "Pantalla"], [
    ["Inicio de sesión", "Correo y contraseña, Iniciar sesión, Olvidé mi contraseña, Registrarse", "2 Iniciar sesión (y 4 y 5 para recuperar)"],
    ["Registro", "Nombre completo, correo, contraseña con confirmación, Registrarse, volver al login", "3 Registrarse"],
    ["Inicio (splash)", "Logo, nombre de la empresa y misión", "1 Inicio"],
    ["CRUD de una tabla", "Buscador, listado con Ver, Editar y Eliminar con confirmación, botón para crear", "22 Vehículos, 23 Ver, 24 Agregar, 25 Editar, 26 Confirmar baja"],
    ["Perfil de usuario", "Datos editables, Guardar cambios, Cerrar sesión", "35 Perfil"],
], [3.4, 7.4, 6.2])
para("En el sistema, eliminar un vehículo es darlo de baja: queda \"Fuera de servicio\" y deja de ofrecerse, pero no se borra, "
     "para no perder el historial de los pedidos que hizo. Por eso el botón dice \"Baja\" y pide confirmación.", negrita="Para la defensa: ")

doc.add_heading("Todas las pantallas", level=2)
ent = re.findall(r"\{ g: '([^']+)', id: '([^']+)', t: '([^']+)', d: '([^']+)'", (BASE / "gallery.js").read_text(encoding="utf-8"))
tabla(["N°", "Rol", "Pantalla", "Qué muestra"], [[str(i), g, t, d] for i, (g, _, t, d) in enumerate(ent, 1)], [1.0, 2.8, 4.6, 8.6], size=8.5)

# ================= dfd =================
doc.add_heading("6. Procesos que faltan en los DFDs", level=1)
para("Salen de comparar cada pantalla con los DFDs actuales. Ya cuentan con que Gestionar Usuarios, Gestionar Vehículo y Asignar recursos se parten como vimos en la guía de lectura.")
tabla(["Proceso que falta", "Pantalla", "Qué hace con los datos"], [
    ["Registrar Cliente", "3 Registrarse", "Guarda al cliente nuevo en Usuarios, con rol Cliente"],
    ["Recuperar Clave", "4 y 5 Recuperar contraseña", "Busca el correo en Usuarios y guarda una clave temporal"],
    ["Cambiar Clave", "6 Cambio obligatorio", "Guarda la clave nueva y apaga el cambio obligatorio"],
    ["Modificar Perfil", "35 Perfil", "Guarda los datos propios del usuario"],
    ["Registrar Inventario", "11 Pedir mudanza, paso 4", "Guarda los objetos en INVENTARIO_DETALLES"],
    ["Consultar Pedidos", "14 y 15 Mis pedidos", "Lee los pedidos del cliente (consulta)"],
    ["Confirmar Pedido", "18 y 19 Detalle del pedido", "Guarda el pedido como confirmado o rechazado"],
    ["Registrar Cobranza", "21 Registrar cobro", "Guarda en COBRANZAS y DETALLE_COBRANZAS"],
    ["Reactivar Usuario", "29 Detalle del usuario", "Vuelve a poner activo a un usuario dado de baja"],
    ["Registrar / Modificar Servicio", "31 Servicios", "Guarda el catálogo y los precios"],
    ["Consultar Mudanzas Asignadas", "32 Mis mudanzas", "Lee los pedidos asignados al empleado (consulta)"],
    ["Actualizar Estado del Pedido", "33 y 34 Detalle de la mudanza", "Guarda \"En curso\" o \"Finalizado\""],
], [4.6, 4.6, 7.8])
doc.add_heading("Lo que la app usa y el DER todavía no tiene", level=2)
for t in [
    "PEDIDOS no tiene columna de estado (Pendiente, Confirmado, Asignado, En curso, Finalizado, Rechazado).",
    "PEDIDOS no tiene direcciones de origen y destino.",
    "ROLES no tiene el rol Cliente.",
    "INVENTARIO_DETALLES no tiene la descripción del objeto (qué es).",
    "No hay dónde guardar qué empleados van a cada mudanza: RECURSOS_ASIGNADOS necesita id_pedido e id_usuario.",
]:
    item(t)
para("La app los usa igual para que la demo funcione. Agregarlos al DER lo decide el grupo.", size=9.5, color=GRIS)

doc.save(OUT)
print("OK", OUT)
